from __future__ import annotations

import argparse
import atexit
import csv
import hashlib
import re
import unicodedata
from calendar import monthrange
from datetime import date
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pywintypes
import win32com.client


_EXCEL_APPLICATION = None


def _excel_application():
    global _EXCEL_APPLICATION
    if _EXCEL_APPLICATION is None:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AutomationSecurity = 3
        _EXCEL_APPLICATION = excel
    return _EXCEL_APPLICATION


def close_excel_application() -> None:
    global _EXCEL_APPLICATION
    if _EXCEL_APPLICATION is not None:
        _EXCEL_APPLICATION.Quit()
        _EXCEL_APPLICATION = None


atexit.register(close_excel_application)


TARGET_RE = re.compile(r"^\d+_(20\d{2})年(\d{1,2})月全国各地区原保险保费收入情况表(?:_|\.xls)")
TITLE_PERIOD_RE = re.compile(r"(20\d{2})年\s*0?(\d{1,2})月")
OUTPUT_NAMES = {
    "facts": "region_premium_facts.parquet",
    "mapping": "region_mapping.csv",
    "quality": "quality_report.csv",
}
OFFICIAL_FALLBACKS = {
    "2024-12": {
        "file": "official_fallbacks/2024-12.xls",
        "url": "https://www.nfra.gov.cn/chinese/docfile/2025/7ce738c5596743c3b4e284f340022cbb.xls",
        "sha256": "d81ca0ae4f2d275c63c5f27518ee182b35fc3098d4f8c845365eb68886e9ffe8",
    },
    "2026-02": {
        "file": "official_fallbacks/2026-02.xls",
        "url": "https://www.nfra.gov.cn/chinese/docfile/2026/1e10b96c585e499890af17ec6bd6060e.xls",
        "sha256": "6dd5cf8ded42ddb939643b6b8e3c600164e66551a62a9f423c5c556f24afc6a2",
    },
}

# Source column, header aliases, canonical product code and product name.
# Every column measures the same thing (original premium income); the old
# cleaner incorrectly modeled the five product breakdowns as five metrics.
METRICS = [
    ("premium_total", "合计", "all", "合计"),
    ("premium_property", "财产保险|财产险", "property_insurance", "财产险"),
    ("premium_life", "寿险", "life_insurance", "寿险"),
    ("premium_accident", "意外险", "accident_insurance", "意外险"),
    ("premium_health", "健康险", "health_insurance", "健康险"),
]

# code, canonical name, type, order, related province code, source aliases
REGIONS = [
    ("CN", "全国", "national", 0, "", ["全国", "全国合计", "全  国"]),
    ("HQ", "公司本级", "company_head_office", 1, "", ["公司本级", "集团、总公司本级"]),
    ("110000", "北京", "province", 10, "", ["北京"]),
    ("120000", "天津", "province", 20, "", ["天津"]),
    ("130000", "河北", "province", 30, "", ["河北"]),
    ("140000", "山西", "province", 40, "", ["山西"]),
    ("150000", "内蒙古", "province", 50, "", ["内蒙古"]),
    ("210000", "辽宁", "province", 60, "", ["辽宁"]),
    ("220000", "吉林", "province", 70, "", ["吉林"]),
    ("230000", "黑龙江", "province", 80, "", ["黑龙江"]),
    ("310000", "上海", "province", 90, "", ["上海"]),
    ("320000", "江苏", "province", 100, "", ["江苏"]),
    ("330000", "浙江", "province", 110, "", ["浙江"]),
    ("340000", "安徽", "province", 120, "", ["安徽"]),
    ("350000", "福建", "province", 130, "", ["福建"]),
    ("360000", "江西", "province", 140, "", ["江西"]),
    ("370000", "山东", "province", 150, "", ["山东"]),
    ("410000", "河南", "province", 160, "", ["河南"]),
    ("420000", "湖北", "province", 170, "", ["湖北"]),
    ("430000", "湖南", "province", 180, "", ["湖南"]),
    ("440000", "广东", "province", 190, "", ["广东"]),
    ("450000", "广西", "province", 200, "", ["广西"]),
    ("460000", "海南", "province", 210, "", ["海南"]),
    ("500000", "重庆", "province", 220, "", ["重庆"]),
    ("510000", "四川", "province", 230, "", ["四川"]),
    ("520000", "贵州", "province", 240, "", ["贵州"]),
    ("530000", "云南", "province", 250, "", ["云南"]),
    ("540000", "西藏", "province", 260, "", ["西藏"]),
    ("610000", "陕西", "province", 270, "", ["陕西"]),
    ("620000", "甘肃", "province", 280, "", ["甘肃"]),
    ("630000", "青海", "province", 290, "", ["青海"]),
    ("640000", "宁夏", "province", 300, "", ["宁夏"]),
    ("650000", "新疆", "province", 310, "", ["新疆"]),
    ("210200", "大连", "planned_city", 320, "210000", ["大连"]),
    ("330200", "宁波", "planned_city", 330, "330000", ["宁波"]),
    ("350200", "厦门", "planned_city", 340, "350000", ["厦门"]),
    ("370200", "青岛", "planned_city", 350, "370000", ["青岛"]),
    ("440300", "深圳", "planned_city", 360, "440000", ["深圳"]),
]


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", str(value))).strip()


REGION_BY_ALIAS = {
    normalize_text(alias): {
        "region_code": code,
        "region_name": name,
        "region_type": kind,
        "region_order": order,
        "parent_region_code": parent,
    }
    for code, name, kind, order, parent, aliases in REGIONS
    for alias in aliases
}


def discover_files(input_dir: Path) -> list[tuple[Path, int, int]]:
    found = []
    for path in input_dir.rglob("*"):
        if not path.is_file() or path.name.startswith("~$") or path.suffix.lower() not in {".xls", ".xlsx"}:
            continue
        match = TARGET_RE.search(path.name)
        if match and 1 <= int(match.group(2)) <= 12:
            found.append((path, int(match.group(1)), int(match.group(2))))
    return sorted(found, key=lambda item: (item[1], item[2], item[0].name))


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def format_excel_value(value: object, number_format: str | None) -> str:
    if value is None:
        return ""
    if not isinstance(value, (int, float, Decimal)):
        return str(value).strip()
    if not number_format or number_format.lower() == "general":
        return str(value).strip()

    number = Decimal(str(value))
    sections = number_format.split(";")
    section = sections[1] if number < 0 and len(sections) > 1 else sections[0]
    percent = "%" in section
    shown = abs(number) * (100 if percent else 1)
    decimals = re.search(r"\.([0#]+)", section)
    places = len(decimals.group(1)) if decimals else 0
    rounded = shown.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    rendered = f"{rounded:,.{places}f}" if "," in section else f"{rounded:.{places}f}"
    if number < 0:
        rendered = f"({rendered})" if "(" in section else f"-{rendered}"
    return rendered + ("%" if percent else "")


def xlsx_sheets(path: Path) -> list[dict]:
    values_wb = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=False)
    formulas_wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=False)
    sheets = []
    try:
        for name in values_wb.sheetnames:
            values_ws = values_wb[name]
            formulas_ws = formulas_wb[name]
            rows = max(values_ws.max_row, formulas_ws.max_row)
            cols = max(values_ws.max_column, formulas_ws.max_column)
            values, display_values, formulas, formats = [], [], [], []
            for row in range(1, rows + 1):
                value_row, display_row, formula_row, format_row = [], [], [], []
                for col in range(1, cols + 1):
                    value_cell = values_ws.cell(row, col)
                    formula_cell = formulas_ws.cell(row, col)
                    raw = formula_cell.value
                    value_row.append(value_cell.value)
                    display_row.append(value_cell.value)
                    formula_row.append(raw if isinstance(raw, str) and raw.startswith("=") else None)
                    format_row.append(formula_cell.number_format or None)
                values.append(value_row)
                display_values.append(display_row)
                formulas.append(formula_row)
                formats.append(format_row)
            sheets.append(
                {
                    "name": name,
                    "values": values,
                    "display_values": display_values,
                    "formulas": formulas,
                    "formats": formats,
                }
            )
    finally:
        values_wb.close()
        formulas_wb.close()
    return sheets


def _range_matrix(raw: object, rows: int, cols: int) -> list[list[object]]:
    if rows == 1 and cols == 1:
        return [[raw]]
    if rows == 1:
        return [list(raw)]
    return [list(row) for row in raw]


def _pad_matrix(matrix: list[list[object]], top: int, left: int, width: int) -> list[list[object]]:
    padded_width = left + width
    return [[None] * padded_width for _ in range(top)] + [
        [None] * left + row + [None] * (width - len(row)) for row in matrix
    ]


def _worksheet_content_bounds(worksheet) -> tuple[int, int, int, int] | None:
    """Return the absolute zero-based bounds of constants and formulas."""
    bounds = None
    for cell_type in (2, -4123):  # xlCellTypeConstants, xlCellTypeFormulas
        try:
            areas = worksheet.Cells.SpecialCells(cell_type).Areas
        except pywintypes.com_error:
            continue
        for area in areas:
            first_row, first_col = int(area.Row) - 1, int(area.Column) - 1
            last_row = first_row + int(area.Rows.Count) - 1
            last_col = first_col + int(area.Columns.Count) - 1
            if bounds is None:
                bounds = first_row, last_row, first_col, last_col
            else:
                top, bottom, left, right = bounds
                bounds = (
                    min(top, first_row),
                    max(bottom, last_row),
                    min(left, first_col),
                    max(right, last_col),
                )
    return bounds


def xls_sheets(path: Path) -> list[dict]:
    """Read legacy XLS once through Excel, retaining stored and rendered values."""
    excel = _excel_application()
    workbook = None
    try:
        workbook = excel.Workbooks.Open(str(path.resolve()), 0, True)
        sheets = []
        for worksheet in workbook.Worksheets:
            bounds = _worksheet_content_bounds(worksheet)
            if bounds is None:
                sheets.append(
                    {
                        "name": str(worksheet.Name),
                        "values": [],
                        "display_values": [],
                        "formulas": [],
                        "formats": [],
                    }
                )
                continue
            first_row, last_row, first_col, last_col = bounds
            top, left = first_row, first_col
            rows = last_row - first_row + 1
            cols = last_col - first_col + 1
            used = worksheet.Range(
                worksheet.Cells.Item(top + 1, left + 1),
                worksheet.Cells.Item(last_row + 1, last_col + 1),
            )
            values = _range_matrix(used.Value2, rows, cols)
            raw_formulas = _range_matrix(used.Formula, rows, cols)
            formulas = [
                [value if isinstance(value, str) and value.startswith("=") else None for value in row]
                for row in raw_formulas
            ]
            formats, display_values = [], []
            for row in range(1, rows + 1):
                format_row, display_row = [], []
                for col in range(1, cols + 1):
                    cell = worksheet.Cells.Item(top + row, left + col)
                    number_format = cell.NumberFormat
                    display = cell.Text
                    if isinstance(display, str) and display.strip("#") == "":
                        display = format_excel_value(
                            values[row - 1][col - 1], number_format
                        )
                    format_row.append(number_format)
                    display_row.append(display)
                formats.append(format_row)
                display_values.append(display_row)
            sheets.append(
                {
                    "name": str(worksheet.Name),
                    "values": _pad_matrix(values, top, left, cols),
                    "display_values": _pad_matrix(display_values, top, left, cols),
                    "formulas": _pad_matrix(formulas, top, left, cols),
                    "formats": _pad_matrix(formats, top, left, cols),
                }
            )
        return sheets
    finally:
        if workbook is not None:
            workbook.Close(False)


def get_cell(sheet: dict, row: int, col: int, field: str = "values"):
    matrix = sheet[field]
    if row < 1 or col < 1 or row > len(matrix) or col > len(matrix[row - 1]):
        return None
    return matrix[row - 1][col - 1]


def find_table(sheets: list[dict]) -> tuple[dict, int, dict[str, int], dict[str, str]]:
    for sheet in sheets:
        for row_number, row in enumerate(sheet["values"], start=1):
            normalized = [normalize_text(value) for value in row]
            if "地区" not in normalized or "合计" not in normalized:
                continue
            positions = {value: index + 1 for index, value in enumerate(normalized) if value}
            property_label = "财产保险" if "财产保险" in positions else "财产险" if "财产险" in positions else ""
            if not property_label or not all(label in positions for label in ("寿险", "意外险", "健康险")):
                continue
            columns = {
                "region": positions["地区"],
                "premium_total": positions["合计"],
                "premium_property": positions[property_label],
                "premium_life": positions["寿险"],
                "premium_accident": positions["意外险"],
                "premium_health": positions["健康险"],
            }
            raw_labels = {
                key: str(get_cell(sheet, row_number, col) or "")
                for key, col in columns.items()
                if key != "region"
            }
            return sheet, row_number, columns, raw_labels
    raise ValueError("未找到包含地区及五个指标的表头")


def decimal_value(value: object) -> tuple[Decimal | None, str | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, None
    if isinstance(value, bool):
        return None, "NON_NUMERIC_VALUE"
    text = unicodedata.normalize("NFKC", str(value)).strip().replace(",", "")
    if text.startswith("#"):
        return None, "EXCEL_ERROR"
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, "NON_NUMERIC_VALUE"
    if number.as_tuple().exponent < -6:
        return number.quantize(Decimal("0.000001")), "PRECISION_ROUNDED_6DP"
    return number, "NEGATIVE_VALUE" if number < 0 else None


def source_cell(row: int, col: int) -> str:
    return f"{openpyxl.utils.get_column_letter(col)}{row}"


def extract_footnotes(sheet: dict, after_row: int) -> str:
    notes, started = [], False
    for row in sheet["values"][after_row:]:
        for value in row:
            if value is None or not str(value).strip():
                continue
            text = str(value).strip()
            if not started and normalize_text(text).startswith("注:"):
                started = True
            if started and text not in notes:
                notes.append(text)
    return "\n".join(notes)


def expected_region_codes(schema_version: str) -> set[str]:
    codes = {code for code, _, kind, _, _, _ in REGIONS if kind != "company_head_office"}
    if schema_version != "V3_2025_PLUS":
        codes.add("HQ")
    return codes


def precision_unit(values: list[Decimal]) -> Decimal:
    return Decimal("1") if values and all(value == value.to_integral_value() for value in values) else Decimal("0.01")


def parse_file(
    path: Path,
    year: int,
    month: int,
    input_dir: Path,
    source_file_label: str | None = None,
    source_url: str = "",
) -> tuple[list[dict], dict]:
    sheets = xlsx_sheets(path) if path.suffix.lower() == ".xlsx" else xls_sheets(path)
    sheet, header_row, columns, raw_metric_labels = find_table(sheets)
    source_format = path.suffix.lower().lstrip(".")

    rows = []
    unknown_regions = []
    last_data_row = header_row
    for row_number in range(header_row + 1, len(sheet["values"]) + 1):
        raw_region = get_cell(sheet, row_number, columns["region"])
        normalized_region = normalize_text(raw_region)
        if normalized_region in REGION_BY_ALIAS:
            rows.append((row_number, raw_region, REGION_BY_ALIAS[normalized_region]))
            last_data_row = row_number
            continue
        row_text = [normalize_text(value) for value in sheet["values"][row_number - 1]]
        if any(text.startswith("注:") for text in row_text if text):
            break
        numeric_like = any(
            get_cell(sheet, row_number, columns[code]) not in (None, "")
            for code, _, _, _ in METRICS
        )
        if normalized_region and numeric_like:
            unknown_regions.append(f"{source_cell(row_number, columns['region'])}:{raw_region}")

    if not rows:
        raise ValueError("表头后未识别到地区数据")

    present_codes = {region["region_code"] for _, _, region in rows}
    schema_version = "V1_2020_2023" if year <= 2023 else "V2_2024" if year == 2024 else "V3_2025_PLUS"
    footnotes = extract_footnotes(sheet, last_data_row)
    accounting_basis = "2026_FINANCIAL_INSTRUMENTS" if "金融工具会计准则" in footnotes else "2009_INSURANCE_ACCOUNTING"
    statistical_scope = (
        "EXCLUDES_RISK_DISPOSAL_INSTITUTIONS"
        if "风险处置" in footnotes
        else "STANDARD_PUBLISHED_SCOPE"
    )
    period_start = date(year, 1, 1)
    period_end = date(year, month, monthrange(year, month)[1])
    release_id = f"{year:04d}-{month:02d}"
    relative_path = source_file_label or path.relative_to(input_dir).as_posix()
    source_origin = "official_download" if source_url else "local_attachment"
    sha256 = file_hash(path)

    pre_header = " ".join(
        str(value) for row in sheet["values"][:header_row] for value in row if value is not None and str(value).strip()
    )
    title_periods = {(int(y), int(m)) for y, m in TITLE_PERIOD_RE.findall(pre_header)}
    title_period_mismatch = bool(title_periods and (year, month) not in title_periods)
    unit_found = "亿元" in normalize_text(pre_header)

    facts = []
    values_by_region: dict[str, dict[str, Decimal | None]] = {}
    formula_cache_missing = 0
    missing_values = 0
    negative_values = 0
    value_errors = 0
    for row_number, raw_region, region in rows:
        region_values = {}
        for product_order, (column_code, _, product_line, product_name) in enumerate(METRICS):
            col = columns[column_code]
            cached = get_cell(sheet, row_number, col)
            display_raw = get_cell(sheet, row_number, col, "display_values")
            formula = get_cell(sheet, row_number, col, "formulas")
            number_format = get_cell(sheet, row_number, col, "formats")
            value_storage, _ = decimal_value(cached)
            value_input = display_raw if source_format == "xls" else cached
            value, flag = decimal_value(value_input)
            value_display = "" if display_raw is None else str(display_raw).strip()
            flags = []
            if formula and value is None:
                flags.append("FORMULA_CACHE_MISSING")
                formula_cache_missing += 1
            if flag:
                flags.append(flag)
                if flag == "NEGATIVE_VALUE":
                    negative_values += 1
                elif flag != "PRECISION_ROUNDED_6DP":
                    value_errors += 1
            if value is None:
                missing_values += 1
            region_values[column_code] = value
            facts.append(
                {
                    "release_id": release_id,
                    "period_start": period_start,
                    "period_end": period_end,
                    "period_basis": "YTD",
                    "region_code": region["region_code"],
                    "region_name": region["region_name"],
                    "region_name_raw": "" if raw_region is None else str(raw_region),
                    "region_type": region["region_type"],
                    "region_order": region["region_order"],
                    "parent_region_code": region["parent_region_code"],
                    "metric_code": "original_premium_income",
                    "metric_name": "原保险保费收入",
                    "metric_name_raw": raw_metric_labels[column_code],
                    "metric_order": 0,
                    "product_line": product_line,
                    "product_line_name": product_name,
                    "product_line_raw": raw_metric_labels[column_code],
                    "product_order": product_order,
                    "metric_path": f"original_premium_income.{product_line}",
                    "value": value,
                    "value_storage": value_storage,
                    "value_display": value_display,
                    "value_decimal": value,
                    "value_raw": value_display,
                    "value_status": "missing" if value is None else "reported_zero" if value == 0 else "reported",
                    "unit": "CNY_100M",
                    "schema_version": schema_version,
                    "scope_version": "HQ_INCLUDED" if schema_version != "V3_2025_PLUS" else "REGION_ONLY",
                    "statistical_scope_version": statistical_scope,
                    "accounting_basis_version": accounting_basis,
                    "source_file": relative_path,
                    "source_file_hash": sha256,
                    "source_format": source_format,
                    "source_origin": source_origin,
                    "source_url": source_url,
                    "source_sheet": sheet["name"],
                    "source_cell": source_cell(row_number, col),
                    "formula_raw": formula,
                    "number_format": number_format,
                    "footnotes": footnotes,
                    "quality_flags": "|".join(flags),
                }
            )
        values_by_region[region["region_code"]] = region_values

    reported_values = [fact["value"] for fact in facts if fact["value"] is not None]
    source_precision_unit = precision_unit(reported_values)
    row_identity_tolerance = Decimal("2.5") if source_precision_unit == 1 else Decimal("0.05")
    row_diffs = []
    for metric_values in values_by_region.values():
        parts = [metric_values.get(code) for code in ("premium_property", "premium_life", "premium_accident", "premium_health")]
        total = metric_values.get("premium_total")
        if total is not None and all(value is not None for value in parts):
            row_diffs.append(abs(total - sum(parts, Decimal(0))))

    national_deltas = {}
    national_rounding_tolerances = {}
    for column_code, _, _, _ in METRICS:
        national_value = values_by_region.get("CN", {}).get(column_code)
        subordinate_values = [
            metrics[column_code]
            for code, metrics in values_by_region.items()
            if code != "CN" and metrics.get(column_code) is not None
        ]
        national_deltas[column_code] = (
            None if national_value is None else national_value - sum(subordinate_values, Decimal(0))
        )
        national_rounding_tolerances[column_code] = source_precision_unit * Decimal(len(subordinate_values) + 1) / 2
    national_delta = national_deltas["premium_total"]
    expected_codes = expected_region_codes(schema_version)
    missing_codes = sorted(expected_codes - present_codes)
    extra_codes = sorted(present_codes - expected_codes)

    qc = {
        "record_type": "FILE",
        "severity": "PASS",
        "period": release_id,
        "source_file": relative_path,
        "source_origin": source_origin,
        "source_url": source_url,
        "source_sheet": sheet["name"],
        "schema_version": schema_version,
        "status": "CLEAN",
        "fact_count": len(facts),
        "region_count": len(rows),
        "expected_region_count": len(expected_codes),
        "row_identity_max_abs_diff": max(row_diffs, default=None),
        "row_identity_tolerance": row_identity_tolerance,
        "national_minus_subregions": national_delta,
        "national_minus_subregions_by_metric": "|".join(
            f"{code}={national_deltas[code]}" for code, _, _, _ in METRICS
        ),
        "numeric_precision": "integer" if source_precision_unit == 1 else "decimal",
        "duplicate_key_count": 0,
        "missing_value_count": missing_values,
        "formula_cache_missing_count": formula_cache_missing,
        "unknown_region_count": len(unknown_regions),
        "negative_value_count": negative_values,
        "cumulative_decrease_count": 0,
        "detail": "",
    }
    details = []
    if missing_codes or extra_codes:
        qc["severity"], qc["status"] = "ERROR", "REGION_SET_MISMATCH"
        details.append(f"missing_regions={missing_codes}; extra_regions={extra_codes}")
    if unknown_regions:
        qc["severity"], qc["status"] = "ERROR", "UNKNOWN_REGION"
        details.append("unknown_regions=" + ",".join(unknown_regions))
    if formula_cache_missing or value_errors:
        qc["severity"], qc["status"] = "ERROR", "VALUE_ERROR"
        details.append(f"formula_cache_missing={formula_cache_missing}; value_errors={value_errors}")
    if missing_values:
        qc["severity"], qc["status"] = "ERROR", "MISSING_VALUE"
        details.append(f"missing_values={missing_values}")
    expected_fact_count = len(expected_codes) * len(METRICS)
    if len(facts) != expected_fact_count:
        qc["severity"], qc["status"] = "ERROR", "FACT_COUNT_MISMATCH"
        details.append(f"fact_count={len(facts)}; expected={expected_fact_count}")
    if title_period_mismatch:
        qc["severity"], qc["status"] = "ERROR", "TITLE_PERIOD_MISMATCH"
        details.append(f"title_periods={sorted(title_periods)}")
    if not unit_found:
        qc["severity"], qc["status"] = "ERROR", "UNIT_NOT_FOUND"
        details.append("unit_亿元_not_found")
    warnings = []
    if row_diffs and max(row_diffs) > row_identity_tolerance:
        warnings.append(f"row_identity_max_abs_diff={max(row_diffs)}")
    for column_code, _, _, _ in METRICS:
        delta = national_deltas[column_code]
        national_value = values_by_region.get("CN", {}).get(column_code)
        if delta is None:
            continue
        relative_delta = abs(delta / national_value) if national_value else Decimal(0)
        if abs(delta) > national_rounding_tolerances[column_code] and (
            abs(delta) > Decimal("1") or relative_delta > Decimal("0.0005")
        ):
            warnings.append(f"national_delta_{column_code}={delta}")
    if negative_values:
        warnings.append(f"negative_values={negative_values}")
    if warnings and qc["severity"] == "PASS":
        qc["severity"], qc["status"] = "WARN", "CHECK_WARNING"
    details.extend(warnings)
    qc["detail"] = "; ".join(details)
    return facts, qc


def missing_periods(periods: set[tuple[int, int]]) -> list[str]:
    if not periods:
        return []
    start = min(year * 12 + month - 1 for year, month in periods)
    end = max(year * 12 + month - 1 for year, month in periods)
    return [
        f"{index // 12:04d}-{index % 12 + 1:02d}"
        for index in range(start, end + 1)
        if (index // 12, index % 12 + 1) not in periods
    ]


def apply_dataset_checks(
    facts: list[dict], quality: list[dict], discovered_periods: set[tuple[int, int]]
) -> tuple[int, list[str]]:
    keys = [
        (row["release_id"], row["region_code"], row["metric_code"], row["product_line"])
        for row in facts
    ]
    duplicate_count = len(keys) - len(set(keys))
    if duplicate_count:
        for row in quality:
            if row["record_type"] == "FILE":
                row["duplicate_key_count"] = duplicate_count
                row["severity"], row["status"] = "ERROR", "DUPLICATE_KEY"

    file_qc = {row["period"]: row for row in quality if row["record_type"] == "FILE"}
    precision_by_release = {
        row["period"]: Decimal("1") if row.get("numeric_precision") == "integer" else Decimal("0.01")
        for row in quality
        if row["record_type"] == "FILE"
    }
    previous: dict[tuple, tuple[str, Decimal, Decimal]] = {}
    for row in sorted(
        facts,
        key=lambda item: (
            item["period_end"],
            item["region_order"],
            item["metric_order"],
            item["product_order"],
        ),
    ):
        if row["value"] is None:
            continue
        key = (
            row["period_end"].year,
            row["region_code"],
            row["metric_code"],
            row["product_line"],
            row["scope_version"],
            row["statistical_scope_version"],
            row["accounting_basis_version"],
        )
        tolerance = precision_by_release.get(row["release_id"], Decimal("0.01"))
        if key in previous and row["value"] + max(tolerance, previous[key][2]) < previous[key][1]:
            qc = file_qc[row["release_id"]]
            qc["cumulative_decrease_count"] += 1
            if qc["severity"] == "PASS":
                qc["severity"], qc["status"] = "WARN", "CUMULATIVE_DECREASE"
            message = (
                f"累计值下降:{row['region_code']}/{row['metric_code']}/"
                f"{row['product_line']} {previous[key][0]}={previous[key][1]} -> "
                f"{row['release_id']}={row['value']}"
            )
            qc["detail"] = "; ".join(filter(None, [qc["detail"], message]))
        previous[key] = (row["release_id"], row["value"], tolerance)

    gaps = missing_periods(discovered_periods)
    for period in gaps:
        quality.append(
            {
                "record_type": "PERIOD",
                "severity": "WARN",
                "period": period,
                "source_file": "",
                "source_origin": "",
                "source_url": "",
                "source_sheet": "",
                "schema_version": "",
                "status": "MISSING_PERIOD",
                "fact_count": "",
                "region_count": "",
                "expected_region_count": "",
                "row_identity_max_abs_diff": "",
                "row_identity_tolerance": "",
                "national_minus_subregions": "",
                "national_minus_subregions_by_metric": "",
                "numeric_precision": "",
                "duplicate_key_count": "",
                "missing_value_count": "",
                "formula_cache_missing_count": "",
                "unknown_region_count": "",
                "negative_value_count": "",
                "cumulative_decrease_count": "",
                "detail": "源文件缺失；不能由相邻累计值推导该月单月值",
            }
        )
    return duplicate_count, gaps


FACT_SCHEMA = pa.schema(
    [
        ("release_id", pa.string()),
        ("period_start", pa.date32()),
        ("period_end", pa.date32()),
        ("period_basis", pa.string()),
        ("region_code", pa.string()),
        ("region_name", pa.string()),
        ("region_name_raw", pa.string()),
        ("region_type", pa.string()),
        ("region_order", pa.int16()),
        ("parent_region_code", pa.string()),
        ("metric_code", pa.string()),
        ("metric_name", pa.string()),
        ("metric_name_raw", pa.string()),
        ("metric_order", pa.int8()),
        ("product_line", pa.string()),
        ("product_line_name", pa.string()),
        ("product_line_raw", pa.string()),
        ("product_order", pa.int8()),
        ("metric_path", pa.string()),
        ("value", pa.decimal128(20, 6)),
        ("value_storage", pa.decimal128(20, 6)),
        ("value_display", pa.string()),
        ("value_decimal", pa.decimal128(20, 6)),
        ("value_raw", pa.string()),
        ("value_status", pa.string()),
        ("unit", pa.string()),
        ("schema_version", pa.string()),
        ("scope_version", pa.string()),
        ("statistical_scope_version", pa.string()),
        ("accounting_basis_version", pa.string()),
        ("source_file", pa.string()),
        ("source_file_hash", pa.string()),
        ("source_format", pa.string()),
        ("source_origin", pa.string()),
        ("source_url", pa.string()),
        ("source_sheet", pa.string()),
        ("source_cell", pa.string()),
        ("formula_raw", pa.string()),
        ("number_format", pa.string()),
        ("footnotes", pa.string()),
        ("quality_flags", pa.string()),
    ]
)


def write_mapping(path: Path) -> None:
    fields = ["alias_raw", "alias_normalized", "region_code", "region_name", "region_type", "region_order", "related_province_code"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for code, name, kind, order, parent, aliases in REGIONS:
            for alias in aliases:
                writer.writerow(
                    {
                        "alias_raw": alias,
                        "alias_normalized": normalize_text(alias),
                        "region_code": code,
                        "region_name": name,
                        "region_type": kind,
                        "region_order": order,
                        "related_province_code": parent,
                    }
                )


def write_quality(path: Path, rows: list[dict]) -> None:
    fields = list(rows[0])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def run(input_dir: Path, output_dir: Path, check_only: bool) -> None:
    files = discover_files(input_dir)
    if not files:
        raise SystemExit("未找到目标 Excel 文件")

    discovered_periods = {(year, month) for _, year, month in files}
    if len(discovered_periods) != len(files):
        raise SystemExit("同一月份发现多个源文件，请先确认应保留的版本")

    all_facts, quality = [], []
    for path, year, month in files:
        release_id = f"{year:04d}-{month:02d}"
        source_file_label = path.relative_to(input_dir).as_posix()
        try:
            facts, qc = parse_file(path, year, month, input_dir)
            all_facts.extend(facts)
            quality.append(qc)
            print(f"OK {year:04d}-{month:02d} {path.name}: {len(facts)} facts")
        except Exception as exc:
            fallback = OFFICIAL_FALLBACKS.get(release_id)
            fallback_path = input_dir / fallback["file"] if fallback else None
            if fallback_path and fallback_path.is_file() and "无法只读打开 xls" in str(exc):
                try:
                    if file_hash(fallback_path) != fallback["sha256"]:
                        raise ValueError("官方备用源哈希不匹配")
                    facts, qc = parse_file(
                        fallback_path,
                        year,
                        month,
                        input_dir,
                        source_file_label=source_file_label,
                        source_url=fallback["url"],
                    )
                    qc["detail"] = "; ".join(filter(None, [qc["detail"], "本地源被占用，使用监管网站官方副本"]))
                    all_facts.extend(facts)
                    quality.append(qc)
                    print(f"FALLBACK {year:04d}-{month:02d} {path.name}: {len(facts)} facts")
                    continue
                except Exception as fallback_exc:
                    exc = RuntimeError(f"本地源读取失败：{exc}; 官方副本读取失败：{fallback_exc}")
            quality.append(
                {
                    "record_type": "FILE",
                    "severity": "ERROR",
                    "period": release_id,
                    "source_file": source_file_label,
                    "source_origin": "local_attachment",
                    "source_url": "",
                    "source_sheet": "",
                    "schema_version": "",
                    "status": "PARSE_ERROR",
                    "fact_count": 0,
                    "region_count": 0,
                    "expected_region_count": "",
                    "row_identity_max_abs_diff": "",
                    "row_identity_tolerance": "",
                    "national_minus_subregions": "",
                    "national_minus_subregions_by_metric": "",
                    "numeric_precision": "",
                    "duplicate_key_count": "",
                    "missing_value_count": "",
                    "formula_cache_missing_count": "",
                    "unknown_region_count": "",
                    "negative_value_count": "",
                    "cumulative_decrease_count": "",
                    "detail": f"{type(exc).__name__}: {exc}",
                }
            )
            print(f"ERROR {year:04d}-{month:02d} {path.name}: {exc}")

    all_facts.sort(
        key=lambda item: (
            item["period_end"],
            item["region_order"],
            item["metric_order"],
            item["product_order"],
        )
    )
    duplicate_count, gaps = apply_dataset_checks(all_facts, quality, discovered_periods)
    parse_errors = [row for row in quality if row["status"] == "PARSE_ERROR"]
    assert not duplicate_count, f"发现 {duplicate_count} 个重复事实键"
    assert len({row["release_id"] for row in all_facts}) + len(parse_errors) == len(files), "文件与发布期数量不一致"
    expected_fact_count = sum((38 if year <= 2024 else 37) * len(METRICS) for _, year, _ in files)
    blocking_errors = [row for row in quality if row["severity"] == "ERROR"]

    if check_only:
        print(
            f"CHECK COMPLETED: files={len(files)}, facts={len(all_facts)}/{expected_fact_count}, "
            f"missing_periods={gaps}, blocking_errors={len(blocking_errors)}"
        )
        if blocking_errors:
            raise SystemExit(1)
        return

    output_dir.mkdir(parents=True, exist_ok=True)
    write_mapping(output_dir / OUTPUT_NAMES["mapping"])
    write_quality(output_dir / OUTPUT_NAMES["quality"], quality)
    if blocking_errors:
        print(f"BLOCKED: {len(blocking_errors)} quality errors; mapping and quality report were written, facts were not published")
        raise SystemExit(1)
    assert len(all_facts) == expected_fact_count, f"事实数 {len(all_facts)} != 预期 {expected_fact_count}"
    pq.write_table(pa.Table.from_pylist(all_facts, schema=FACT_SCHEMA), output_dir / OUTPUT_NAMES["facts"], compression="zstd")
    print(f"WROTE {len(all_facts)} facts to {output_dir}")


def main() -> None:
    parser = argparse.ArgumentParser(description="清洗全国各地区原保险保费收入 Excel 为可追溯长表")
    parser.add_argument("--input-dir", type=Path, default=Path.cwd())
    parser.add_argument("--output-dir", type=Path, default=Path.cwd() / "output_region_premium_clean")
    parser.add_argument("--check-only", action="store_true", help="全量解析并校验，不写结果文件")
    args = parser.parse_args()
    run(args.input_dir.resolve(), args.output_dir.resolve(), args.check_only)


if __name__ == "__main__":
    main()
