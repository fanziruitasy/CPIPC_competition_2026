from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

import pdfplumber
from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from pypdf import PdfReader

try:
    import fitz  # type: ignore
except Exception:  # pragma: no cover - optional dependency
    fitz = None


SUPPORTED_EXTS = {".pdf", ".docx", ".doc"}
DEFAULT_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(chunk_size), b""):
            h.update(chunk)
    return h.hexdigest()


def stable_doc_id(path: Path) -> str:
    m = re.match(r"^(\d+)_", path.name)
    if m:
        return f"doc_{m.group(1)}"
    return "doc_" + hashlib.sha1(path.name.encode("utf-8")).hexdigest()[:10]


def clean_text(text: str) -> str:
    text = text.replace("\u0000", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t\u3000]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def one_line(text: str, max_len: int = 160) -> str:
    text = re.sub(r"\s+", " ", text or "").strip()
    return text[:max_len]


def normalize_for_compare(text: str) -> str:
    text = clean_text(text)
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，。；：、“”‘’（）()\[\]【】《》,.，;:!?！？\-—_]", "", text)
    return text


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_json(path: Path, obj: Any) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def append_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    with path.open("a", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def md_escape_cell(value: Any) -> str:
    text = "" if value is None else str(value)
    text = clean_text(text).replace("\n", "<br>")
    return text.replace("|", "\\|")


def table_to_markdown(rows: list[list[Any]]) -> str:
    if not rows:
        return ""
    width = max((len(r) for r in rows), default=0)
    norm = [(r + [""] * (width - len(r)))[:width] for r in rows]
    header = [md_escape_cell(c) for c in norm[0]]
    sep = ["---"] * width
    body = [[md_escape_cell(c) for c in r] for r in norm[1:]]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(sep) + " |"]
    lines.extend("| " + " | ".join(r) + " |" for r in body)
    return "\n".join(lines)


def normalize_table_matrix(table: list[list[Any]]) -> list[list[str]]:
    width = max((len(r) for r in table if r), default=0)
    out = []
    for row in table:
        normalized = []
        for cell in (row or [])[:width]:
            normalized.append(clean_text("" if cell is None else str(cell)))
        if len(normalized) < width:
            normalized.extend([""] * (width - len(normalized)))
        out.append(normalized)
    return out


def table_cells(matrix: list[list[str]]) -> list[dict[str, Any]]:
    cells = []
    for ridx, row in enumerate(matrix):
        for cidx, text in enumerate(row):
            cells.append({
                "row_index": ridx,
                "column_index": cidx,
                "text": text,
                "is_empty": not bool(text),
            })
    return cells


def flatten_form_kv(matrix: list[list[str]]) -> list[dict[str, Any]]:
    pairs = []
    seen = set()
    for ridx, row in enumerate(matrix):
        non_empty = [(cidx, text) for cidx, text in enumerate(row) if text]
        if len(non_empty) < 2:
            continue
        for idx in range(0, len(non_empty) - 1, 2):
            key_col, key = non_empty[idx]
            value_col, value = non_empty[idx + 1]
            if key == value or len(key) > 80:
                continue
            sig = (key, value, ridx, key_col)
            if sig in seen:
                continue
            seen.add(sig)
            pairs.append({
                "key": key,
                "value": value,
                "row_index": ridx,
                "key_column_index": key_col,
                "value_column_index": value_col,
            })
    return pairs[:200]


def enrich_table_record(table: list[list[Any]], include_full_content: bool = True) -> dict[str, Any]:
    matrix = normalize_table_matrix(table)
    summary = summarize_table_cells(matrix)
    markdown = table_to_markdown(matrix)
    summary.update({
        "is_preview_truncated": len(" ".join(" ".join(r) for r in matrix)) > len(summary["text_preview"]),
        "sample_rows_truncated": len(matrix) > len(summary["sample_rows"]),
        "table_markdown": markdown,
    })
    if include_full_content:
        summary.update({
            "raw_rows": matrix,
            "normalized_rows": matrix,
            "cells": table_cells(matrix),
            "flattened_kv": flatten_form_kv(matrix),
        })
    return summary


def flatten_complex_header(matrix: list[list[str]], max_header_rows: int = 3) -> dict[str, Any]:
    if not matrix:
        return {"header_rows": 0, "flattened_columns": [], "data_rows": []}
    width = max((len(r) for r in matrix), default=0)
    if width == 0:
        return {"header_rows": 0, "flattened_columns": [], "data_rows": []}
    header_rows = 1
    limit = min(max_header_rows, len(matrix))
    for ridx in range(1, limit):
        row = matrix[ridx]
        non_empty_values = [c for c in row if c]
        numeric_values = [c for c in non_empty_values if re.fullmatch(r"[-+]?\d+(\.\d+)?%?|[-+]?\d+(\.\d+)?[~-][-+]?\d+(\.\d+)?%?", c)]
        first_cell = row[0] if row else ""
        data_like_first_cell = bool(re.fullmatch(r"\d+|[（(]?[一二三四五六七八九十]+[）)]?", first_cell))
        if non_empty_values and (len(numeric_values) >= max(2, len(non_empty_values) // 2) or data_like_first_cell):
            break
        empty_ratio = sum(1 for c in row if not c) / max(1, width)
        shortish = sum(1 for c in row if c and len(c) <= 30)
        if empty_ratio >= 0.3 or shortish >= max(1, width // 2):
            header_rows = ridx + 1
        else:
            break
    flattened = []
    for cidx in range(width):
        parts = []
        for ridx in range(header_rows):
            value = matrix[ridx][cidx] if cidx < len(matrix[ridx]) else ""
            if value and value not in parts:
                parts.append(value)
        flattened.append("/".join(parts) if parts else f"column_{cidx + 1}")
    return {
        "header_rows": header_rows,
        "flattened_columns": flattened,
        "data_rows": matrix[header_rows:],
    }


def summarize_table_cells(table: list[list[Any]], max_rows: int = 5, max_cols: int = 8) -> dict[str, Any]:
    row_count = len(table)
    col_count = max((len(r) for r in table if r), default=0)
    non_empty = 0
    text_parts = []
    sample_rows = []
    for ridx, row in enumerate(table):
        sample = []
        for cidx, cell in enumerate(row or []):
            text = clean_text("" if cell is None else str(cell))
            if text:
                non_empty += 1
                if len(text_parts) < 24:
                    text_parts.append(text)
            if ridx < max_rows and cidx < max_cols:
                sample.append(text)
        if ridx < max_rows:
            sample_rows.append(sample)
    return {
        "row_count": row_count,
        "column_count": col_count,
        "non_empty_cell_count": non_empty,
        "text_preview": one_line(" ".join(text_parts), 240),
        "sample_rows": sample_rows,
    }


def infer_heading_level(text: str, style_name: str | None = None) -> int | None:
    style_name = style_name or ""
    m = re.search(r"Heading\s*(\d+)|标题\s*(\d+)", style_name, flags=re.I)
    if m:
        return int(next(g for g in m.groups() if g))
    s = text.strip()
    if not s:
        return None
    if re.match(r"^附件\s*\d+", s):
        return 1
    if re.match(r"^第[一二三四五六七八九十百]+章", s):
        return 1
    if re.match(r"^第[一二三四五六七八九十百]+节", s):
        return 2
    if re.match(r"^[一二三四五六七八九十]+、", s):
        return 2
    if re.match(r"^（[一二三四五六七八九十]+）", s):
        return 3
    if re.match(r"^\d+[.．、]", s):
        return 4
    return None


def page_role(text_len: int, table_count: int, image_count: int = 0) -> str:
    if text_len == 0 and image_count == 0:
        return "blank"
    if text_len < 80 and image_count > 0:
        return "decorative_or_image"
    if table_count > 0 and text_len < 300:
        return "table"
    if table_count > 0:
        return "mixed"
    return "text"

# docx底层遍历，交替产出paragraph、Table，解决docx里表格和段落互相穿插的遍历问题
def iter_block_items(parent: DocxDocumentType):
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def extract_docx_notes_and_objects(path: Path) -> dict[str, Any]:
    out = {
        "footnote_count": 0,
        "endnote_count": 0,
        "media_count": 0,
        "ole_object_count": 0,
        "formula_object_count": 0,
        "has_external_relationships": False,
    }
    if not zipfile.is_zipfile(path):
        return out
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        out["media_count"] = sum(1 for n in names if n.startswith("word/media/") and not n.endswith("/"))
        out["ole_object_count"] = sum(1 for n in names if n.startswith("word/embeddings/") and not n.endswith("/"))
        for note_file, key in [
            ("word/footnotes.xml", "footnote_count"),
            ("word/endnotes.xml", "endnote_count"),
        ]:
            if note_file in names:
                root = etree.fromstring(z.read(note_file))
                tag = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                notes = root.findall(f".//{tag}{'footnote' if 'foot' in note_file else 'endnote'}")
                out[key] = len([n for n in notes if n.get(f"{tag}type") not in {"separator", "continuationSeparator"}])
        for n in names:
            if n.endswith(".rels") and n.startswith("word/_rels/"):
                xml = etree.fromstring(z.read(n))
                for rel in xml:
                    if rel.get("TargetMode") == "External":
                        out["has_external_relationships"] = True
        out["formula_object_count"] = out["ole_object_count"]
    return out


def extract_docx_object_inventory(path: Path, doc_id: str) -> list[dict[str, Any]]:
    if not zipfile.is_zipfile(path):
        return []
    inventory = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        for idx, name in enumerate([n for n in names if n.startswith("word/embeddings/") and not n.endswith("/")], 1):
            suffix = Path(name).suffix.lower()
            inventory.append({
                "object_id": f"obj_{doc_id}_{idx:04d}",
                "doc_id": doc_id,
                "source_type": "word",
                "object_type": "ole_embedding",
                "package_path": name,
                "file_extension": suffix,
                "size_bytes": len(z.read(name)),
                "semantic_status": "inventory_only",
                "semantic_restore_priority": "high" if suffix in {".xlsx", ".xls", ".docx", ".doc"} else "medium",
                "semantic_text": None,
                "anchor": None,
                "context_text": None,
                "quality_flags": ["needs_semantic_restore"],
            })
        media_start = len(inventory)
        for idx, name in enumerate([n for n in names if n.startswith("word/media/") and not n.endswith("/")], 1):
            suffix = Path(name).suffix.lower()
            inventory.append({
                "object_id": f"obj_{doc_id}_{media_start + idx:04d}",
                "doc_id": doc_id,
                "source_type": "word",
                "object_type": "media",
                "package_path": name,
                "file_extension": suffix,
                "size_bytes": len(z.read(name)),
                "semantic_status": "inventory_only",
                "semantic_restore_priority": "medium" if suffix in {".png", ".jpg", ".jpeg", ".emf", ".wmf"} else "low",
                "semantic_text": None,
                "anchor": None,
                "context_text": None,
                "quality_flags": ["may_need_ocr_or_formula_restore"],
            })
    return inventory


@dataclass
class Paths:
    input_dir: Path
    output_dir: Path
    soffice: Path
    markdown_dir: Path = field(init=False)
    metadata_dir: Path = field(init=False)
    intermediate_dir: Path = field(init=False)
    converted_dir: Path = field(init=False)
    reports_dir: Path = field(init=False)

    def __post_init__(self) -> None:
        self.markdown_dir = self.output_dir / "markdown"
        self.metadata_dir = self.output_dir / "metadata"
        self.intermediate_dir = self.output_dir / "intermediate"
        self.converted_dir = self.output_dir / "converted"
        self.reports_dir = self.output_dir / "reports"


class IngestionPipeline:
    def __init__(self, paths: Paths, timeout_seconds: int = 120) -> None:
        self.paths = paths
        self.timeout_seconds = timeout_seconds
        self.manifest_rows: list[dict[str, Any]] = []
        self.evidence_rows: list[dict[str, Any]] = []
        self.table_rows: list[dict[str, Any]] = []
        self.pdf_enhanced_table_rows: list[dict[str, Any]] = []
        self.pdf_page_rows: list[dict[str, Any]] = []
        self.word_block_rows: list[dict[str, Any]] = []
        self.object_rows: list[dict[str, Any]] = []
        self.errors: list[dict[str, Any]] = []
        self.document_rows: list[dict[str, Any]] = []

    def run(self) -> None:
        for directory in [
            self.paths.output_dir,
            self.paths.markdown_dir,
            self.paths.metadata_dir,
            self.paths.intermediate_dir,
            self.paths.converted_dir,
            self.paths.reports_dir,
        ]:
            ensure_dir(directory)
        files = self.discover_files()
        for idx, source in enumerate(files, 1):
            print(f"[{idx}/{len(files)}] {source.name}", flush=True)
            try:
                self.process_one(source)
            except Exception as exc:
                self.record_error(source, "file_failed", exc)
        self.write_outputs()

    def discover_files(self) -> list[Path]:
        files = []
        for p in self.paths.input_dir.iterdir():
            if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS and not p.name.startswith("~$"):
                files.append(p)
        return sorted(files, key=lambda x: (x.suffix.lower(), x.name))

    def source_record(self, path: Path, parse_strategy: str) -> dict[str, Any]:
        return {
            "doc_id": stable_doc_id(path),
            "source_path": str(path),
            "source_filename": path.name,
            "source_hash": "sha256:" + sha256_file(path),
            "file_type": path.suffix.lower().lstrip("."),
            "size_bytes": path.stat().st_size,
            "discovered_at": now_iso(),
            "parse_strategy": parse_strategy,
            "parse_status": "started",
            "parser_versions": self.parser_versions(),
        }

    def parser_versions(self) -> dict[str, Any]:
        return {
            "python": sys.version.split()[0],
            "pdfplumber": getattr(pdfplumber, "__version__", None),
            "pymupdf_available": fitz is not None,
            "soffice_path": str(self.paths.soffice),
        }

    def process_one(self, source: Path) -> None:
        ext = source.suffix.lower()
        strategy = "dual_path" if ext == ".doc" else "direct"
        rec = self.source_record(source, strategy)
        self.manifest_rows.append(rec)
        if ext == ".pdf":
            meta = self.parse_pdf(source, rec)
        elif ext == ".docx":
            meta = self.parse_docx(source, rec, source_role="source_docx")
        elif ext == ".doc":
            meta = self.parse_doc(source, rec)
        else:
            raise ValueError(f"Unsupported extension: {ext}")
        rec["parse_status"] = meta.get("parse_status", "success")
        rec["quality_flags"] = meta.get("quality_flags", [])
        self.document_rows.append({
            "doc_id": rec["doc_id"],
            "source_filename": source.name,
            "file_type": rec["file_type"],
            "document_title": meta.get("document_title"),
            "markdown_path": str(self.paths.markdown_dir / f"{rec['doc_id']}.md"),
            "metadata_path": str(self.paths.metadata_dir / f"{rec['doc_id']}.json"),
            "parse_status": rec["parse_status"],
            "quality_flags": rec["quality_flags"],
        })

    def parse_pdf(self, source: Path, rec: dict[str, Any]) -> dict[str, Any]:
        doc_id = rec["doc_id"]
        pymupdf_pages = self.inspect_pdf_with_pymupdf(source)
        metadata: dict[str, Any] = {
            **rec,
            "parse_status": "success",
            "document_title": source.stem,
            "pages": [],
            "tables": [],
            "evidence_units": [],
            "quality_flags": [],
        }
        reader = PdfReader(str(source))
        metadata["page_count"] = len(reader.pages)
        metadata["pdf_outlines_count"] = self.count_pdf_outlines(reader)
        metadata["pymupdf_audit"] = {
            "available": fitz is not None,
            "page_count": len(pymupdf_pages),
        }
        md_lines = self.front_matter(metadata, "pdf")
        md_lines.append(f"# {source.stem}\n")
        with pdfplumber.open(str(source)) as pdf:
            for pidx, page in enumerate(pdf.pages, 1):
                raw_text = page.extract_text() or ""
                text = clean_text(raw_text)
                tables = page.extract_tables() or []
                image_count = len(getattr(page, "images", []) or [])
                role = page_role(len(text), len(tables), image_count)
                page_rec = {
                    "doc_id": doc_id,
                    "page_number": pidx,
                    "width": page.width,
                    "height": page.height,
                    "rotation": getattr(page, "rotation", 0),
                    "text_length": len(text),
                    "char_count": len(page.chars),
                    "table_candidate_count": len(tables),
                    "image_count": image_count,
                    "role": role,
                    "ocr_decision": "not_required" if len(text) > 0 else "review_required",
                    "quality_flags": [],
                }
                if pidx <= len(pymupdf_pages):
                    page_rec["pymupdf"] = pymupdf_pages[pidx - 1]
                    image_count = max(image_count, pymupdf_pages[pidx - 1].get("image_count", 0))
                    page_rec["image_count"] = image_count
                if "\x00" in raw_text:
                    page_rec["quality_flags"].append("null_char_removed")
                if len(text) < 80:
                    page_rec["quality_flags"].append("low_text_page")
                self.pdf_page_rows.append(page_rec)
                metadata["pages"].append(page_rec)
                md_lines.append(f"\n<!-- page: {pidx} role: {role} -->\n")
                if text:
                    ev_id = f"ev_{doc_id}_p{pidx:03d}_text"
                    md_lines.append(f'<a id="{ev_id}"></a>\n')
                    md_lines.append(text + "\n")
                    ev = {
                        "evidence_id": ev_id,
                        "doc_id": doc_id,
                        "source_type": "pdf",
                        "content_type": "page_text",
                        "text": text,
                        "location": {"page_number": pidx, "bbox": None},
                        "markdown_anchor": ev_id,
                        "quality_flags": page_rec["quality_flags"],
                    }
                    self.evidence_rows.append(ev)
                    metadata["evidence_units"].append(ev_id)
                for tidx, table in enumerate(tables, 1):
                    table_id = f"table_{doc_id}_p{pidx:03d}_{tidx:03d}"
                    accepted, reason = self.accept_pdf_table(table)
                    table_summary = enrich_table_record(table, include_full_content=accepted)
                    header_info = flatten_complex_header(table_summary.get("normalized_rows", []))
                    rejection_class = None if accepted else self.classify_rejected_pdf_table(table, text, role)
                    trow = {
                        "table_id": table_id,
                        "doc_id": doc_id,
                        "source_type": "pdf",
                        "page_number": pidx,
                        "row_count": table_summary["row_count"],
                        "column_count": table_summary["column_count"],
                        "non_empty_cell_count": table_summary["non_empty_cell_count"],
                        "text_preview": table_summary["text_preview"],
                        "sample_rows": table_summary["sample_rows"],
                        "is_preview_truncated": table_summary["is_preview_truncated"],
                        "sample_rows_truncated": table_summary["sample_rows_truncated"],
                        "raw_rows": table_summary.get("raw_rows"),
                        "normalized_rows": table_summary.get("normalized_rows"),
                        "cells": table_summary.get("cells"),
                        "flattened_kv": table_summary.get("flattened_kv"),
                        "table_markdown": table_summary["table_markdown"] if accepted else None,
                        "complex_header": header_info if accepted else None,
                        "decision": "accepted" if accepted else "rejected",
                        "decision_reason": reason,
                        "rejection_class": rejection_class,
                        "cross_page_group_id": None,
                    }
                    self.table_rows.append(trow)
                    metadata["tables"].append(trow)
                    if accepted:
                        md_lines.append(f"\n<!-- table: {table_id} page: {pidx} -->\n")
                        md_lines.append(f"### 表格 {table_id}\n\n")
                        md_lines.append(table_summary["table_markdown"] + "\n")
                        ev_id = f"ev_{table_id}"
                        self.evidence_rows.append({
                            "evidence_id": ev_id,
                            "doc_id": doc_id,
                            "source_type": "pdf",
                            "content_type": "table",
                            "text": table_summary["table_markdown"],
                            "location": {"page_number": pidx, "table_id": table_id},
                            "markdown_anchor": table_id,
                            "quality_flags": [],
                        })
                        metadata["evidence_units"].append(ev_id)
        self.enhance_pdf_table_groups(metadata)
        metadata["quality_summary"] = self.summarize_pdf(metadata)
        self.write_doc_outputs(doc_id, md_lines, metadata)
        return metadata

    def enhance_pdf_table_groups(self, metadata: dict[str, Any]) -> None:
        accepted = [t for t in metadata.get("tables", []) if t.get("source_type") == "pdf" and t.get("decision") == "accepted"]
        accepted.sort(key=lambda t: (t.get("page_number") or 0, t.get("table_id") or ""))
        group_index = 0
        previous: dict[str, Any] | None = None
        for table in accepted:
            if previous and self.is_cross_page_table_candidate(previous, table):
                if not previous.get("cross_page_group_id"):
                    group_index += 1
                    previous["cross_page_group_id"] = f"cpt_{metadata['doc_id']}_{group_index:04d}"
                table["cross_page_group_id"] = previous["cross_page_group_id"]
                table["cross_page_relation"] = "continuation_candidate"
            else:
                table["cross_page_relation"] = "single_or_group_start"
            previous = table
        for table in accepted:
            self.pdf_enhanced_table_rows.append({
                "table_id": table.get("table_id"),
                "doc_id": table.get("doc_id"),
                "page_number": table.get("page_number"),
                "row_count": table.get("row_count"),
                "column_count": table.get("column_count"),
                "non_empty_cell_count": table.get("non_empty_cell_count"),
                "cross_page_group_id": table.get("cross_page_group_id"),
                "cross_page_relation": table.get("cross_page_relation"),
                "complex_header": table.get("complex_header"),
                "flattened_columns": (table.get("complex_header") or {}).get("flattened_columns", []),
                "normalized_rows": table.get("normalized_rows"),
                "table_markdown": table.get("table_markdown"),
            })

    def is_cross_page_table_candidate(self, left: dict[str, Any], right: dict[str, Any]) -> bool:
        if (right.get("page_number") or 0) != (left.get("page_number") or 0) + 1:
            return False
        left_cols = left.get("column_count") or 0
        right_cols = right.get("column_count") or 0
        if not left_cols or not right_cols or abs(left_cols - right_cols) > 1:
            return False
        left_header = ((left.get("complex_header") or {}).get("flattened_columns") or [])[: min(left_cols, right_cols)]
        right_header = ((right.get("complex_header") or {}).get("flattened_columns") or [])[: min(left_cols, right_cols)]
        if left_header and right_header:
            ratio = SequenceMatcher(None, " ".join(left_header), " ".join(right_header)).ratio()
            if ratio >= 0.45:
                return True
        left_preview = left.get("text_preview") or ""
        right_preview = right.get("text_preview") or ""
        return SequenceMatcher(None, left_preview[:160], right_preview[:160]).ratio() >= 0.55

    def inspect_pdf_with_pymupdf(self, source: Path) -> list[dict[str, Any]]:
        if fitz is None:
            return []
        pages: list[dict[str, Any]] = []
        with fitz.open(str(source)) as doc:
            for idx, page in enumerate(doc, 1):
                rect = page.rect
                try:
                    text_blocks = page.get_text("blocks")
                except Exception:
                    text_blocks = []
                pages.append({
                    "page_number": idx,
                    "width": float(rect.width),
                    "height": float(rect.height),
                    "rotation": int(page.rotation),
                    "image_count": len(page.get_images(full=True)),
                    "text_block_count": len(text_blocks),
                })
        return pages

    def count_pdf_outlines(self, reader: PdfReader) -> int:
        def walk(items: Any) -> int:
            total = 0
            if not isinstance(items, list):
                return 0
            for item in items:
                if isinstance(item, list):
                    total += walk(item)
                else:
                    total += 1
            return total
        try:
            return walk(reader.outline)
        except Exception:
            return 0

    def accept_pdf_table(self, table: list[list[Any]]) -> tuple[bool, str]:
        summary = summarize_table_cells(table)
        rows = summary["row_count"]
        cols = summary["column_count"]
        non_empty = summary["non_empty_cell_count"]
        if rows < 2 or cols < 2:
            return False, "too_small"
        if non_empty < 4:
            return False, "mostly_empty"
        return True, "grid_and_text_present"

    def classify_rejected_pdf_table(self, table: list[list[Any]], page_text: str, page_role_value: str) -> str:
        summary = summarize_table_cells(table)
        preview = summary["text_preview"]
        rows = summary["row_count"]
        cols = summary["column_count"]
        non_empty = summary["non_empty_cell_count"]
        context = preview + " " + one_line(page_text, 300)
        form_keywords = ["名称", "日期", "签字", "签章", "盖章", "填报", "联系人", "电话", "金额", "账号", "项目"]
        table_keywords = ["表", "项目", "指标", "合计", "资产", "负债", "资本", "风险", "余额"]
        if non_empty == 0:
            return "blank_grid_or_layout_artifact"
        if rows < 2 or cols < 2:
            if any(k in context for k in form_keywords + table_keywords):
                return "single_axis_form_fragment"
            return "layout_or_caption_noise"
        if non_empty < 4 and any(k in context for k in form_keywords):
            return "blank_form"
        if non_empty < 4:
            return "mostly_empty_grid"
        if page_role_value == "decorative_or_image":
            return "image_border_misdetected"
        return "needs_manual_review"

    def parse_docx(self, source: Path, rec: dict[str, Any], source_role: str) -> dict[str, Any]:
        doc_id = rec["doc_id"]
        doc = DocxDocument(str(source))
        notes = extract_docx_notes_and_objects(source)
        object_inventory = extract_docx_object_inventory(source, doc_id)
        metadata: dict[str, Any] = {
            **rec,
            "parse_status": "success",
            "document_title": source.stem,
            "source_role": source_role,
            "paragraph_count": 0,
            "table_count": len(doc.tables),
            "blocks": [],
            "tables": [],
            "evidence_units": [],
            "notes_and_objects": notes,
            "object_inventory": object_inventory,
            "quality_flags": [],
        }
        self.object_rows.extend(object_inventory)
        md_lines = self.front_matter(metadata, "word")
        md_lines.append(f"# {source.stem}\n")
        table_seen = 0
        para_seen = 0
        for order, block in enumerate(iter_block_items(doc), 1):
            if isinstance(block, Paragraph):
                text = clean_text(block.text)
                if not text:
                    continue
                para_seen += 1
                level = infer_heading_level(text, getattr(block.style, "name", None))
                block_id = f"block_{doc_id}_{order:05d}"
                block_rec = {
                    "block_id": block_id,
                    "doc_id": doc_id,
                    "source_type": "word",
                    "order": order,
                    "block_type": "heading" if level else "paragraph",
                    "heading_level": level,
                    "text": text,
                    "style_name": getattr(block.style, "name", None),
                }
                self.word_block_rows.append(block_rec)
                metadata["blocks"].append(block_rec)
                if level:
                    md_lines.append(f"\n{'#' * min(level + 1, 6)} {text}\n")
                else:
                    ev_id = f"ev_{doc_id}_b{order:05d}"
                    md_lines.append(f'\n<a id="{ev_id}"></a>\n{text}\n')
                    self.evidence_rows.append({
                        "evidence_id": ev_id,
                        "doc_id": doc_id,
                        "source_type": "word",
                        "content_type": "paragraph",
                        "text": text,
                        "location": {"block_order": order, "paragraph_index": para_seen},
                        "markdown_anchor": ev_id,
                        "quality_flags": [],
                    })
            elif isinstance(block, Table):
                table_seen += 1
                table_id = f"table_{doc_id}_{table_seen:04d}"
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                table_summary = enrich_table_record(rows, include_full_content=True)
                table_rec = {
                    "table_id": table_id,
                    "doc_id": doc_id,
                    "source_type": "word",
                    "order": order,
                    "row_count": table_summary["row_count"],
                    "column_count": table_summary["column_count"],
                    "non_empty_cell_count": table_summary["non_empty_cell_count"],
                    "text_preview": table_summary["text_preview"],
                    "sample_rows": table_summary["sample_rows"],
                    "is_preview_truncated": table_summary["is_preview_truncated"],
                    "sample_rows_truncated": table_summary["sample_rows_truncated"],
                    "raw_rows": table_summary["raw_rows"],
                    "normalized_rows": table_summary["normalized_rows"],
                    "cells": table_summary["cells"],
                    "flattened_kv": table_summary["flattened_kv"],
                    "table_markdown": table_summary["table_markdown"],
                    "decision": "accepted",
                    "decision_reason": "word_table_in_body",
                }
                self.table_rows.append(table_rec)
                metadata["tables"].append(table_rec)
                block_rec = {
                    "block_id": f"block_{doc_id}_{order:05d}",
                    "doc_id": doc_id,
                    "source_type": "word",
                    "order": order,
                    "block_type": "table",
                    "table_id": table_id,
                    "text": one_line(" ".join(" ".join(map(str, r)) for r in rows)),
                }
                self.word_block_rows.append(block_rec)
                metadata["blocks"].append(block_rec)
                md_lines.append(f"\n<!-- table: {table_id} order: {order} -->\n")
                md_lines.append(f"### 表格 {table_id}\n\n")
                md_lines.append(table_summary["table_markdown"] + "\n")
                ev_id = f"ev_{table_id}"
                self.evidence_rows.append({
                    "evidence_id": ev_id,
                    "doc_id": doc_id,
                    "source_type": "word",
                    "content_type": "table",
                    "text": table_summary["table_markdown"],
                    "location": {"block_order": order, "table_id": table_id},
                    "markdown_anchor": table_id,
                    "quality_flags": [],
                })
                metadata["evidence_units"].append(ev_id)
        metadata["paragraph_count"] = para_seen
        if table_seen != len(doc.tables):
            metadata["quality_flags"].append("table_count_order_scan_mismatch")
        if notes["ole_object_count"]:
            metadata["quality_flags"].append("contains_ole_objects")
        if notes["footnote_count"] or notes["endnote_count"]:
            metadata["quality_flags"].append("contains_notes")
        metadata["quality_summary"] = self.summarize_word(metadata)
        self.write_doc_outputs(doc_id, md_lines, metadata)
        return metadata

    def parse_doc(self, source: Path, rec: dict[str, Any]) -> dict[str, Any]:
        doc_id = rec["doc_id"]
        stem = source.stem
        doc_out = self.paths.converted_dir / "doc" / doc_id
        ensure_dir(doc_out)
        representations = []
        converted_docx = self.convert_with_soffice(source, doc_out, "docx")
        converted_pdf = self.convert_with_soffice(source, doc_out, "pdf")
        converted_txt = self.convert_with_soffice(source, doc_out, "txt:Text")
        for role, path in [
            ("converted_docx", converted_docx),
            ("visual_check_pdf", converted_pdf),
            ("text_baseline", converted_txt),
        ]:
            representations.append({
                "role": role,
                "path": str(path) if path else None,
                "status": "success" if path and path.exists() else "failed",
            })
        if not converted_docx:
            meta = {
                **rec,
                "parse_status": "failed",
                "document_title": stem,
                "representations": representations,
                "quality_flags": ["docx_conversion_failed"],
            }
            write_json(self.paths.metadata_dir / f"{doc_id}.json", meta)
            (self.paths.markdown_dir / f"{doc_id}.md").write_text(
                "\n".join(self.front_matter(meta, "word")) + f"\n# {stem}\n\nDOC conversion failed.\n",
                encoding="utf-8",
            )
            return meta
        docx_rec = {**rec, "source_path": str(converted_docx), "file_type": "docx"}
        meta = self.parse_docx(converted_docx, docx_rec, source_role="converted_from_doc")
        meta["source_path"] = str(source)
        meta["file_type"] = "doc"
        meta["parse_strategy"] = "dual_path"
        meta["representations"] = representations
        meta["selected_representation_role"] = "converted_docx"
        meta["doc_text_consistency"] = self.compare_doc_text_baseline(converted_txt, meta)
        meta["quality_flags"] = sorted(set(meta.get("quality_flags", []) + ["doc_converted_via_libreoffice"]))
        if meta["doc_text_consistency"].get("coverage_ratio", 1.0) < 0.9:
            meta["quality_flags"] = sorted(set(meta["quality_flags"] + ["low_doc_text_baseline_coverage"]))
        meta["quality_summary"] = self.summarize_word(meta)
        md_path = self.paths.markdown_dir / f"{doc_id}.md"
        if md_path.exists():
            md = md_path.read_text(encoding="utf-8")
            md = re.sub(r"source_file: .+", f"source_file: {source.name}", md, count=1)
            md = re.sub(r"file_type: .+", "file_type: doc", md, count=1)
            md_path.write_text(md, encoding="utf-8")
        write_json(self.paths.metadata_dir / f"{doc_id}.json", meta)
        return meta

    def compare_doc_text_baseline(self, txt_path: Path | None, docx_meta: dict[str, Any]) -> dict[str, Any]:
        docx_text = "\n".join(
            block.get("text", "")
            for block in docx_meta.get("blocks", [])
            if block.get("block_type") in {"paragraph", "heading", "table"}
        )
        docx_norm = normalize_for_compare(docx_text)
        if not txt_path or not txt_path.exists():
            return {
                "status": "missing_txt_baseline",
                "baseline_length": 0,
                "docx_length": len(docx_norm),
                "coverage_ratio": None,
                "similarity_ratio": None,
            }
        raw = txt_path.read_bytes()
        baseline_text = None
        encodings = ["utf-8-sig", "gb18030", "latin1"]
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            encodings.insert(0, "utf-16")
        for enc in encodings:
            try:
                baseline_text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if baseline_text is None:
            baseline_text = raw.decode("utf-8", errors="ignore")
        baseline_norm = normalize_for_compare(baseline_text)
        if not baseline_norm or not docx_norm:
            coverage = 0.0 if docx_norm else 1.0
            similarity = 0.0
        else:
            matcher = SequenceMatcher(None, baseline_norm, docx_norm, autojunk=False)
            matched = sum(size for _, _, size in matcher.get_matching_blocks())
            coverage = matched / max(len(baseline_norm), 1)
            similarity = matcher.ratio()
        return {
            "status": "compared",
            "baseline_path": str(txt_path),
            "baseline_length": len(baseline_norm),
            "docx_length": len(docx_norm),
            "coverage_ratio": round(coverage, 4),
            "similarity_ratio": round(similarity, 4),
            "length_delta_ratio": round(abs(len(docx_norm) - len(baseline_norm)) / max(len(baseline_norm), 1), 4),
        }

    def convert_with_soffice(self, source: Path, out_dir: Path, target: str) -> Path | None:
        ensure_dir(out_dir)
        if not self.paths.soffice.exists():
            self.errors.append({
                "source_path": str(source),
                "error_type": "soffice_missing",
                "message": str(self.paths.soffice),
                "created_at": now_iso(),
            })
            return None
        profile = self.paths.output_dir / "tmp" / "lo_profile"
        ensure_dir(profile)
        before = {p.name for p in out_dir.iterdir()} if out_dir.exists() else set()
        cmd = [
            str(self.paths.soffice),
            "--headless",
            "--nologo",
            "--nofirststartwizard",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            target,
            "--outdir",
            str(out_dir),
            str(source),
        ]
        try:
            completed = subprocess.run(
                cmd,
                cwd=str(out_dir),
                capture_output=True,
                text=True,
                timeout=self.timeout_seconds,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            self.errors.append({
                "source_path": str(source),
                "error_type": "soffice_timeout",
                "target": target,
                "message": str(exc),
                "created_at": now_iso(),
            })
            return None
        after = list(out_dir.iterdir())
        new_files = [p for p in after if p.name not in before]
        suffix = ".txt" if target.startswith("txt") else "." + target.split(":")[0]
        candidates = [p for p in new_files if p.suffix.lower() == suffix]
        if not candidates:
            candidates = sorted(out_dir.glob(source.stem + suffix))
        if completed.returncode != 0 or not candidates:
            self.errors.append({
                "source_path": str(source),
                "error_type": "soffice_conversion_failed",
                "target": target,
                "returncode": completed.returncode,
                "stdout": completed.stdout[-1000:],
                "stderr": completed.stderr[-1000:],
                "created_at": now_iso(),
            })
            return None
        return max(candidates, key=lambda p: p.stat().st_mtime)

    def front_matter(self, metadata: dict[str, Any], kind: str) -> list[str]:
        return [
            "---",
            f"doc_id: {metadata['doc_id']}",
            f"source_file: {metadata['source_filename']}",
            f"file_type: {metadata['file_type']}",
            f"document_kind: {kind}",
            f"source_hash: {metadata['source_hash']}",
            f"parse_strategy: {metadata['parse_strategy']}",
            "---",
            "",
        ]

    def summarize_pdf(self, metadata: dict[str, Any]) -> dict[str, Any]:
        pages = metadata.get("pages", [])
        tables = metadata.get("tables", [])
        flags = sorted({f for p in pages for f in p.get("quality_flags", [])})
        metadata["quality_flags"] = flags
        return {
            "page_count": len(pages),
            "low_text_pages": sum(1 for p in pages if "low_text_page" in p.get("quality_flags", [])),
            "table_candidate_count": len(tables),
            "accepted_table_count": sum(1 for t in tables if t.get("decision") == "accepted"),
            "rejected_table_count": sum(1 for t in tables if t.get("decision") == "rejected"),
            "quality_flags": flags,
        }

    def summarize_word(self, metadata: dict[str, Any]) -> dict[str, Any]:
        return {
            "paragraph_count": metadata.get("paragraph_count", 0),
            "table_count": metadata.get("table_count", 0),
            "block_count": len(metadata.get("blocks", [])),
            "notes_and_objects": metadata.get("notes_and_objects", {}),
            "quality_flags": metadata.get("quality_flags", []),
        }

    def write_doc_outputs(self, doc_id: str, md_lines: list[str], metadata: dict[str, Any]) -> None:
        (self.paths.markdown_dir / f"{doc_id}.md").write_text("\n".join(md_lines), encoding="utf-8")
        write_json(self.paths.metadata_dir / f"{doc_id}.json", metadata)

    def record_error(self, source: Path, error_type: str, exc: Exception) -> None:
        self.errors.append({
            "source_path": str(source),
            "error_type": error_type,
            "message": str(exc),
            "traceback": traceback.format_exc(limit=5),
            "created_at": now_iso(),
        })

    def write_outputs(self) -> None:
        for path in [
            self.paths.output_dir / "manifest.jsonl",
            self.paths.output_dir / "documents.jsonl",
            self.paths.output_dir / "evidence_units.jsonl",
            self.paths.intermediate_dir / "structured_tables.jsonl",
            self.paths.intermediate_dir / "pdf_enhanced_tables.jsonl",
            self.paths.intermediate_dir / "pdf_pages.jsonl",
            self.paths.intermediate_dir / "word_blocks.jsonl",
            self.paths.intermediate_dir / "object_inventory.jsonl",
            self.paths.reports_dir / "processing_errors.jsonl",
        ]:
            if path.exists():
                path.unlink()
        append_jsonl(self.paths.output_dir / "manifest.jsonl", self.manifest_rows)
        append_jsonl(self.paths.output_dir / "documents.jsonl", self.document_rows)
        append_jsonl(self.paths.output_dir / "evidence_units.jsonl", self.evidence_rows)
        append_jsonl(self.paths.intermediate_dir / "structured_tables.jsonl", self.table_rows)
        append_jsonl(self.paths.intermediate_dir / "pdf_enhanced_tables.jsonl", self.pdf_enhanced_table_rows)
        append_jsonl(self.paths.intermediate_dir / "pdf_pages.jsonl", self.pdf_page_rows)
        append_jsonl(self.paths.intermediate_dir / "word_blocks.jsonl", self.word_block_rows)
        append_jsonl(self.paths.intermediate_dir / "object_inventory.jsonl", self.object_rows)
        append_jsonl(self.paths.reports_dir / "processing_errors.jsonl", self.errors)
        write_json(self.paths.reports_dir / "quality_report.json", self.quality_report())
        write_json(self.paths.reports_dir / "pdf_quality_report.json", self.pdf_quality_report())
        write_json(self.paths.reports_dir / "word_quality_report.json", self.word_quality_report())

    def quality_report(self) -> dict[str, Any]:
        return {
            "created_at": now_iso(),
            "input_dir": str(self.paths.input_dir),
            "output_dir": str(self.paths.output_dir),
            "manifest_count": len(self.manifest_rows),
            "documents_count": len(self.document_rows),
            "markdown_count": len(list(self.paths.markdown_dir.glob("*.md"))),
            "metadata_count": len(list(self.paths.metadata_dir.glob("*.json"))),
            "evidence_units_count": len(self.evidence_rows),
            "structured_tables_count": len(self.table_rows),
            "pdf_enhanced_tables_count": len(self.pdf_enhanced_table_rows),
            "object_inventory_count": len(self.object_rows),
            "errors_count": len(self.errors),
            "by_type": self.count_by("file_type"),
            "by_status": self.count_by("parse_status"),
        }

    def pdf_quality_report(self) -> dict[str, Any]:
        return {
            "pdf_count": sum(1 for r in self.manifest_rows if r["file_type"] == "pdf"),
            "page_count": len(self.pdf_page_rows),
            "low_text_pages": sum(1 for r in self.pdf_page_rows if "low_text_page" in r.get("quality_flags", [])),
            "table_candidates": sum(r.get("table_candidate_count", 0) for r in self.pdf_page_rows),
            "accepted_tables": sum(1 for r in self.table_rows if r.get("source_type") == "pdf" and r.get("decision") == "accepted"),
            "rejected_tables": sum(1 for r in self.table_rows if r.get("source_type") == "pdf" and r.get("decision") == "rejected"),
            "enhanced_tables": len(self.pdf_enhanced_table_rows),
            "cross_page_table_groups": len({r.get("cross_page_group_id") for r in self.pdf_enhanced_table_rows if r.get("cross_page_group_id")}),
            "complex_header_tables": sum(1 for r in self.pdf_enhanced_table_rows if (r.get("complex_header") or {}).get("header_rows", 0) > 1),
            "rejected_table_classes": self.count_pdf_rejection_classes(),
            "pymupdf_available": fitz is not None,
        }

    def count_pdf_rejection_classes(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for row in self.table_rows:
            if row.get("source_type") == "pdf" and row.get("decision") == "rejected":
                value = str(row.get("rejection_class") or "unclassified")
                out[value] = out.get(value, 0) + 1
        return dict(sorted(out.items(), key=lambda item: item[0]))

    def word_quality_report(self) -> dict[str, Any]:
        word_docs = [r for r in self.manifest_rows if r["file_type"] in {"doc", "docx"}]
        doc_recs = [r for r in self.manifest_rows if r["file_type"] == "doc"]
        docx_recs = [r for r in self.manifest_rows if r["file_type"] == "docx"]
        metas = []
        for p in self.paths.metadata_dir.glob("doc_*.json"):
            try:
                meta = json.loads(p.read_text(encoding="utf-8"))
            except Exception:
                continue
            if meta.get("file_type") in {"doc", "docx"}:
                metas.append(meta)
        return {
            "word_count": len(word_docs),
            "doc_count": len(doc_recs),
            "docx_count": len(docx_recs),
            "paragraph_count": sum(m.get("paragraph_count", 0) for m in metas),
            "table_count": sum(m.get("table_count", 0) for m in metas),
            "footnote_count": sum(m.get("notes_and_objects", {}).get("footnote_count", 0) for m in metas),
            "endnote_count": sum(m.get("notes_and_objects", {}).get("endnote_count", 0) for m in metas),
            "ole_object_count": sum(m.get("notes_and_objects", {}).get("ole_object_count", 0) for m in metas),
            "object_inventory_count": len(self.object_rows),
            "semantic_restore_needed": sum(1 for r in self.object_rows if r.get("semantic_status") == "inventory_only"),
            "doc_converted_success": sum(
                1
                for m in metas
                if m.get("file_type") == "doc"
                and any(r.get("role") == "converted_docx" and r.get("status") == "success" for r in m.get("representations", []))
            ),
            "doc_text_consistency": self.doc_text_consistency_summary(metas),
        }

    def doc_text_consistency_summary(self, metas: list[dict[str, Any]]) -> dict[str, Any]:
        doc_metas = [m for m in metas if m.get("file_type") == "doc"]
        compared = [
            m.get("doc_text_consistency", {})
            for m in doc_metas
            if m.get("doc_text_consistency", {}).get("status") == "compared"
        ]
        ratios = [r.get("coverage_ratio") for r in compared if isinstance(r.get("coverage_ratio"), (int, float))]
        low = [
            {
                "doc_id": m.get("doc_id"),
                "source_filename": m.get("source_filename"),
                "coverage_ratio": m.get("doc_text_consistency", {}).get("coverage_ratio"),
                "similarity_ratio": m.get("doc_text_consistency", {}).get("similarity_ratio"),
            }
            for m in doc_metas
            if isinstance(m.get("doc_text_consistency", {}).get("coverage_ratio"), (int, float))
            and m.get("doc_text_consistency", {}).get("coverage_ratio") < 0.9
        ]
        return {
            "compared_count": len(compared),
            "min_coverage_ratio": round(min(ratios), 4) if ratios else None,
            "avg_coverage_ratio": round(sum(ratios) / len(ratios), 4) if ratios else None,
            "low_coverage_docs": low,
        }

    def count_by(self, key: str) -> dict[str, int]:
        out: dict[str, int] = {}
        for row in self.manifest_rows:
            value = str(row.get(key))
            out[value] = out.get(value, 0) + 1
        return out


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Parse PDF/DOC/DOCX into Markdown + metadata JSON.")
    parser.add_argument("--input-dir", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--soffice", default=DEFAULT_SOFFICE)
    parser.add_argument("--timeout-seconds", type=int, default=120)
    args = parser.parse_args(argv)
    paths = Paths(Path(args.input_dir), Path(args.output_dir), Path(args.soffice))
    pipeline = IngestionPipeline(paths, timeout_seconds=args.timeout_seconds)
    started = time.time()
    pipeline.run()
    print(f"Done in {time.time() - started:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
