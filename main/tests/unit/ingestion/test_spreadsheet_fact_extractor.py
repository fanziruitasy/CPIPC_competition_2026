"""验证专用 Excel 清洗结果到统一事实、证据和 Parquet 的映射。"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq

from trusted_rag.domain.enums import FormulaCacheStatus, QualityStatus
from trusted_rag.ingestion.source_registry import SourceRegistrar
from trusted_rag.ingestion.spreadsheets.classification import WorkbookClassification
from trusted_rag.ingestion.spreadsheets.cleaners import regulatory_tables
from trusted_rag.ingestion.spreadsheets.fact_extractor import (
    UnifiedSpreadsheetExtractor,
    write_facts_parquet,
)


def test_extract_regional_facts_with_cell_evidence_and_formula_review(
    tmp_path: Path,
) -> None:
    """事实和证据必须稳定对应原单元格，缺少公式缓存值时进入人工复核。"""
    source_root = tmp_path / "sources"
    source_root.mkdir()
    workbook_path = source_root / "001_2024年3月全国各地区原保险保费收入情况表.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "原保险保费"
    sheet.append(["2024年3月全国各地区原保险保费收入情况表（单位：亿元）"])
    sheet.append(["地区", "合计", "财产保险", "寿险", "意外险", "健康险"])
    sheet.append(["全国", "=SUM(C3:F3)", 20, 50, 10, 20])
    workbook.save(workbook_path)
    workbook.close()
    source = SourceRegistrar(source_root).register(
        workbook_path,
        knowledge_base_id="kb_excel",
        run_id="registration-run",
    )

    result = UnifiedSpreadsheetExtractor(source_root).extract(
        source,
        run_id="excel-run",
    )

    assert len(result.facts) == 5
    assert len(result.evidence) == 8
    total = result.facts[0]
    assert total.metric_code == "premium_total"
    assert total.period_end == date(2024, 3, 31)
    assert total.unit == "CNY_100M"
    assert total.location.sheet_name == "原保险保费"
    assert total.location.cell_range == "B3"
    assert total.formula == "=SUM(C3:F3)"
    assert total.formula_cache_status is FormulaCacheStatus.MISSING
    assert total.quality.status is QualityStatus.REQUIRES_REVIEW
    total_evidence = next(
        evidence
        for evidence in result.evidence
        if evidence.evidence_id == total.evidence_id
    )
    assert total_evidence.location == total.location
    assert result.document.quality.requires_manual_review is True


def test_write_unified_facts_parquet_preserves_contract_fields(tmp_path: Path) -> None:
    """统一事实写入 Parquet 后保留规范数值、公式状态和证据定位。"""
    source_root = tmp_path / "sources"
    source_root.mkdir()
    workbook_path = source_root / "001_2024年3月全国各地区原保险保费收入情况表.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "原保险保费"
    sheet.append(["2024年3月全国各地区原保险保费收入情况表（单位：亿元）"])
    sheet.append(["地区", "合计", "财产保险", "寿险", "意外险", "健康险"])
    sheet.append(["全国", 100, 20, 50, 10, 20])
    workbook.save(workbook_path)
    workbook.close()
    source = SourceRegistrar(source_root).register(
        workbook_path,
        knowledge_base_id="kb_excel",
        run_id="registration-run",
    )
    result = UnifiedSpreadsheetExtractor(source_root).extract(source, run_id="excel-run")
    output_path = write_facts_parquet(result.facts, tmp_path / "facts.parquet")

    restored = pq.read_table(output_path).to_pylist()

    assert len(restored) == 5
    assert restored[0]["normalized_value"] == "100"
    assert restored[0]["sheet_name"] == "原保险保费"
    assert restored[0]["cell_range"] == "B3"
    assert restored[0]["formula_cache_status"] == "not_formula"
    assert restored[0]["evidence_id"] == result.facts[0].evidence_id


def test_funds_parser_receives_original_file_name(
    tmp_path: Path,
    monkeypatch: object,
) -> None:
    """转换后的稳定文件名不能覆盖资金运用表的原始季度语义。"""
    source_root = tmp_path / "sources"
    source_root.mkdir()
    workbook_path = source_root / "source_stable.xlsx"
    workbook = openpyxl.Workbook()
    workbook.save(workbook_path)
    workbook.close()
    source = SourceRegistrar(source_root).register(
        workbook_path,
        knowledge_base_id="kb_excel",
        run_id="registration-run",
    ).model_copy(
        update={"original_file_name": "017_2024年一季度保险资金运用情况表.xls"}
    )
    captured: dict[str, str] = {}

    def parser(
        path: Path,
        input_dir: Path,
        dataset_id: str,
        year: int,
        source_file_name: str,
    ) -> list[dict[str, object]]:
        del path, input_dir, dataset_id, year
        captured["source_file_name"] = source_file_name
        return []

    monkeypatch.setattr(regulatory_tables, "parse_funds", parser)  # type: ignore[union-attr]
    classification = WorkbookClassification(
        dataset_id="insurance_funds_investment",
        role="fact_table",
        rule_id="test",
        year=2024,
    )

    UnifiedSpreadsheetExtractor(source_root)._extract_raw_facts(
        workbook_path,
        source,
        classification,
    )

    assert captured["source_file_name"] == "017_2024年一季度保险资金运用情况表.xls"
