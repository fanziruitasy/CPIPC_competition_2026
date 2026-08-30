"""将官方 QA 工作簿规范化为可追溯的版本化评测记录。"""

from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path
from typing import Annotated, Any, Literal, cast

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import Field

from trusted_rag.domain.common import ContractModel
from trusted_rag.infrastructure.artifacts import sha256_file, write_json_atomic

EXPECTED_COLUMNS = (
    "id",
    "source_type",
    "difficulty",
    "difficulty_cn",
    "qa_type",
    "question",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "answer",
    "answer_text",
    "evidence",
    "source_title",
    "file_label",
)


class EvaluationRecord(ContractModel):
    """一条可下钻到官方工作簿行的选择题评测记录。"""

    schema_version: Literal["evaluation_record.v1"] = "evaluation_record.v1"
    question_id: Annotated[str, Field(pattern=r"^Q\d{3,}$")]
    source_type: Literal["excel", "word", "pdf"]
    difficulty: Literal["easy", "medium", "hard"]
    difficulty_cn: str
    qa_type: str
    question: Annotated[str, Field(min_length=1)]
    options: dict[Literal["A", "B", "C", "D"], str]
    reference_answer: Literal["A", "B", "C", "D"]
    reference_answer_text: str
    evidence_reference: str
    source_title: str
    file_label: str
    source_sheet: str
    source_row: Annotated[int, Field(ge=2)]


def normalize_official_qa(
    *,
    workbook_path: Path,
    output_root: Path,
    run_id: str,
    expected_question_count: int = 300,
) -> dict[str, Any]:
    """只读加载官方 QA，写出 JSONL 记录和可复现清单。

    :param workbook_path: 官方 QA XLSX 文件。
    :param output_root: 必须尚不存在的评测数据运行目录。
    :param run_id: 版本化评测数据运行标识。
    :param expected_question_count: 期望题目数，用于防止静默漏行。
    :return: 已写入 ``run.json`` 的数据集清单。
    :raises ValueError: 表头、题号、记录数或字段不符合契约。
    :raises FileExistsError: 输出运行目录已经存在。
    """
    if output_root.exists():
        raise FileExistsError(output_root)
    digest_before = sha256_file(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if len(workbook.worksheets) != 1:
        raise ValueError("官方 QA 工作簿必须且只能包含一个工作表。")
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = tuple(str(value).strip() if value is not None else "" for value in next(rows))
    if header != EXPECTED_COLUMNS:
        raise ValueError(f"官方 QA 表头与契约不一致：{header}")
    records: list[EvaluationRecord] = []
    for row_number, values in enumerate(rows, start=2):
        if all(value is None for value in values):
            continue
        item = dict(zip(header, values, strict=True))
        records.append(_record(item, sheet.title, row_number))
    workbook.close()
    if len(records) != expected_question_count:
        raise ValueError(f"官方 QA 应有 {expected_question_count} 题，实际 {len(records)} 题。")
    question_ids = [record.question_id for record in records]
    if len(set(question_ids)) != len(question_ids):
        raise ValueError("官方 QA 题号不唯一。")
    expected_ids = [f"Q{index:03d}" for index in range(1, expected_question_count + 1)]
    if question_ids != expected_ids:
        raise ValueError("官方 QA 题号必须按 Q001 起连续递增。")
    if sha256_file(workbook_path) != digest_before:
        raise RuntimeError("官方 QA 工作簿在规范化过程中发生变化。")

    output_root.mkdir(parents=True, exist_ok=False)
    records_path = output_root / "records.jsonl"
    with records_path.open("w", encoding="utf-8", newline="\n") as stream:
        for record in records:
            stream.write(json.dumps(record.model_dump(mode="json"), ensure_ascii=False) + "\n")
    counts = {
        "source_type": _counts(record.source_type for record in records),
        "difficulty": _counts(record.difficulty for record in records),
        "qa_type": _counts(record.qa_type for record in records),
    }
    manifest = {
        "schema_version": "evaluation_dataset.v1",
        "run_id": run_id,
        "source_workbook_sha256": digest_before,
        "source_sheet": sheet.title,
        "question_count": len(records),
        "question_id_first": records[0].question_id,
        "question_id_last": records[-1].question_id,
        "counts": counts,
        "records_sha256": sha256_file(records_path),
    }
    write_json_atomic(output_root / "run.json", manifest)
    return manifest


def _record(item: dict[str, Any], sheet_name: str, row_number: int) -> EvaluationRecord:
    question_id = _text(item["id"])
    if not re.fullmatch(r"Q\d{3,}", question_id):
        raise ValueError(f"第 {row_number} 行题号不合法：{question_id}")
    return EvaluationRecord(
        question_id=question_id,
        source_type=cast(
            Literal["excel", "word", "pdf"],
            _enum(item["source_type"], row_number, "source_type", ("excel", "word", "pdf")),
        ),
        difficulty=cast(
            Literal["easy", "medium", "hard"],
            _enum(item["difficulty"], row_number, "difficulty", ("easy", "medium", "hard")),
        ),
        difficulty_cn=_required(item["difficulty_cn"], row_number, "difficulty_cn"),
        qa_type=_required(item["qa_type"], row_number, "qa_type"),
        question=_required(item["question"], row_number, "question"),
        options={
            "A": _required(item["option_a"], row_number, "option_a"),
            "B": _required(item["option_b"], row_number, "option_b"),
            "C": _required(item["option_c"], row_number, "option_c"),
            "D": _required(item["option_d"], row_number, "option_d"),
        },
        reference_answer=cast(
            Literal["A", "B", "C", "D"],
            _enum(item["answer"], row_number, "answer", ("A", "B", "C", "D"), upper=True),
        ),
        reference_answer_text=_required(item["answer_text"], row_number, "answer_text"),
        evidence_reference=_required(item["evidence"], row_number, "evidence"),
        source_title=_required(item["source_title"], row_number, "source_title"),
        file_label=_required(item["file_label"], row_number, "file_label"),
        source_sheet=sheet_name,
        source_row=row_number,
    )


def _required(value: Any, row_number: int, field: str) -> str:
    text = _text(value)
    if not text:
        raise ValueError(f"第 {row_number} 行字段 {field} 不能为空。")
    return text


def _enum(
    value: Any,
    row_number: int,
    field: str,
    allowed: tuple[str, ...],
    *,
    upper: bool = False,
) -> str:
    text = _text(value)
    if upper:
        text = text.upper()
    if text not in allowed:
        raise ValueError(f"第 {row_number} 行字段 {field} 不在允许值 {allowed} 中。")
    return text


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _counts(values: Any) -> dict[str, int]:
    counts: dict[str, int] = {}
    for value in values:
        key = str(value)
        counts[key] = counts.get(key, 0) + 1
    return dict(sorted(counts.items()))
